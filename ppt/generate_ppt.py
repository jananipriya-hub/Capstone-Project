from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)  # 16:9 Widescreen

# Colors
BG_COLOR = RGBColor(7, 26, 51)  # Deep Dark Navy (#071A33)
TITLE_COLOR = RGBColor(34, 211, 238)  # Bright Cyan (#22D3EE)
TEXT_COLOR = RGBColor(230, 241, 255)  # Light Blue/White (#E6F1FF)
MUTED_COLOR = RGBColor(140, 168, 198)  # Muted Gray/Blue


def set_slide_background(slide):
  background = slide.background
  fill = background.fill
  fill.solid()
  fill.fore_color.rgb = BG_COLOR


def add_title(slide, text):
  title_box = slide.shapes.add_textbox(
      Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0)
  )
  tf = title_box.text_frame
  tf.word_wrap = True
  p = tf.paragraphs[0]
  p.text = text
  p.font.name = "Segoe UI"
  p.font.size = Pt(32)
  p.font.bold = True
  p.font.color.rgb = TITLE_COLOR


def add_bullets(slide, points, left=0.8, top=1.8, width=11.7, height=5.0):
  box = slide.shapes.add_textbox(
      Inches(left), Inches(top), Inches(width), Inches(height)
  )
  tf = box.text_frame
  tf.word_wrap = True
  for i, pt in enumerate(points):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = "• " + pt
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT_COLOR
    p.space_after = Pt(14)


blank_layout = prs.slide_layouts[6]

# --- SLIDE 1: Title Slide ---
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)
box = slide1.shapes.add_textbox(
    Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.5)
  )
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Credit Risk Analysis & Loan Underwriting System"
p.font.name = "Segoe UI"
p.font.size = Pt(38)
p.font.bold = True
p.font.color.rgb = TITLE_COLOR
p.space_after = Pt(10)

p2 = tf.add_paragraph()
p2.text = (
    "Machine Learning-Powered Decision Support Platform\nAuthor: Janani Priya"
    " N. | B.Tech Information Technology"
)
p2.font.name = "Segoe UI"
p2.font.size = Pt(20)
p2.font.color.rgb = MUTED_COLOR

# --- SLIDE 2: Problem Statement ---
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_title(slide2, "The Challenge of Modern Lending")
add_bullets(
    slide2,
    [
        (
            "Financial institutions face substantial monetary exposure and risk"
            " from high loan default rates."
        ),
        (
            "Traditional manual underwriting processes are slow, inconsistent,"
            " and prone to human oversight."
        ),
        (
            "Objective: Build an automated, intelligent classification"
            " platform that evaluates loan applicants in real-time."
        ),
        (
            "Calculates precise default probabilities and explains"
            " underwriting decisions transparently."
        ),
    ],
)

# --- SLIDE 3: Data & Preprocessing ---
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_title(slide3, "Data Engineering & Preparation")
add_bullets(
    slide3,
    [
        (
            "Source Data: Cleaned credit risk dataset featuring applicant"
            " demographics, financial history, and loan metrics."
        ),
        (
            "Feature Engineering: Mapped ordinal loan grades (A–G), created"
            " Debt-to-Income (DTI) ratios, and extracted credit history depth."
        ),
        (
            "Handling Imbalance: Applied class_weight='balanced' during model"
            " training to ensure proper detection of minority default classes."
        ),
    ],
)

# --- SLIDE 4: EDA Insights ---
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_title(slide4, "Key Risk Drivers & Data Insights")
add_bullets(
    slide4,
    [
        (
            "Income Impact: Lower annual income brackets show a higher"
            " concentration of defaults."
        ),
        (
            "Loan Intent Risk: Debt consolidation and home improvement loans"
            " exhibit higher default rates."
        ),
        (
            "Credit Depth: Shorter credit history lengths strongly correlate"
            " with increased default vulnerability."
        ),
    ],
    width=7.0,
)
# Placeholder text box or note for image insertion on slide 4
img_box = slide4.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.5)
)
img_box.fill.solid()
img_box.fill.fore_color.rgb = RGBColor(14, 42, 77)
img_box.line.color.rgb = RGBColor(28, 68, 116)
tf_img = img_box.text_frame
tf_img.word_wrap = True
p_img = tf_img.paragraphs[0]
p_img.text = "\n\n[ Insert eda_income_patterns_dark.png here ]"
p_img.font.size = Pt(14)
p_img.font.color.rgb = MUTED_COLOR
p_img.alignment = PP_ALIGN.CENTER

# --- SLIDE 5: ML Modeling ---
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_title(slide5, "Predictive Modeling: Random Forest Classifier")
add_bullets(
    slide5,
    [
        (
            "Algorithm Selection: Random Forest chosen for high resistance to"
            " overfitting and capability to capture non-linear interactions."
        ),
        (
            "Evaluation Metrics: Achieved an exceptional ROC-AUC score of"
            " ~0.93, proving powerful discrimination capability."
        ),
        (
            "Model Artifacts: Serialized model (credit_risk_rf_model.pkl),"
            " scalers, and feature lists saved for production integration."
        ),
    ],
)

# --- SLIDE 6: System Architecture ---
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_title(slide6, "Production Deployment & Interactive Interface")
add_bullets(
    slide6,
    [
        (
            "Tech Stack: Built with Python, Scikit-Learn, Streamlit, Pandas,"
            " and Seaborn; version controlled via GitHub."
        ),
        (
            "UI Features: Modular sidebar navigation, secure modal user"
            " login, and persistent session audit logging."
        ),
        (
            "Explainability ('Why This Score'): Generates automated"
            " explanations detailing DTI ratios, loan grades, and credit"
            " depth."
        ),
    ],
    width=7.0,
)
img_box2 = slide6.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.5)
)
img_box2.fill.solid()
img_box2.fill.fore_color.rgb = RGBColor(14, 42, 77)
img_box2.line.color.rgb = RGBColor(28, 68, 116)
tf_img2 = img_box2.text_frame
tf_img2.word_wrap = True
p_img2 = tf_img2.paragraphs[0]
p_img2.text = (
    "\n\n[ Insert Streamlit Input & Prediction Screenshot / Chart here ]"
)
p_img2.font.size = Pt(14)
p_img2.font.color.rgb = MUTED_COLOR
p_img2.alignment = PP_ALIGN.CENTER

# --- SLIDE 7: Conclusion ---
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_title(slide7, "Impact & Future Enhancements")
add_bullets(
    slide7,
    [
        (
            "Business Value: Streamlines institutional decision-making, cuts"
            " down manual review time, and provides transparent audit trails."
        ),
        (
            "Future Scope: Integrate live credit bureau APIs for automated"
            " background pulls and expand cloud monitoring pipelines."
        ),
        (
            "Conclusion: Successfully delivered a fully functional,"
            " theme-adaptive, production-ready machine learning web app."
        ),
    ],
)

prs.save("Credit_Risk_Capstone_Presentation.pptx")
print("Presentation saved successfully as 'Credit_Risk_Capstone_Presentation.pptx'!")